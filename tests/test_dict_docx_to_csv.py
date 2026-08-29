from __future__ import annotations

import base64
import csv
import io
import unittest

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches
from PIL import Image

from dict_docx_to_csv import ConversionError, convert_dictionary_docx_to_csv


def make_png() -> io.BytesIO:
    stream = io.BytesIO()
    Image.new("RGB", (12, 8), (35, 104, 75)).save(stream, "PNG")
    stream.seek(0)
    return stream


def make_docx(include_image: bool, extra_table: bool = False) -> io.BytesIO:
    document = Document()

    if extra_table:
        decoy = document.add_table(rows=2, cols=2)
        decoy.cell(0, 0).text = "Name"
        decoy.cell(0, 1).text = "Value"

    column_count = 3 if include_image else 2
    table = document.add_table(rows=3, cols=column_count)
    headers = [" Word ", "DEFINITION"] + (["Image"] if include_image else [])
    for index, header in enumerate(headers):
        table.cell(0, index).text = header

    table.cell(1, 0).text = "abate"
    paragraph = table.cell(1, 1).paragraphs[0]
    rich_run = paragraph.add_run("become less intense")
    rich_run.bold = True
    rich_run.font.highlight_color = None

    table.cell(2, 0).text = "brisk"
    table.cell(2, 1).text = "active\nand energetic"

    if include_image:
        table.cell(1, 2).paragraphs[0].add_run().add_picture(make_png(), width=Inches(0.25))

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def decode_tsv(content: bytes):
    text = content.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text), delimiter="\t"))


class ConverterTests(unittest.TestCase):
    def test_converts_required_columns_without_image(self):
        result = convert_dictionary_docx_to_csv(
            make_docx(False, extra_table=True),
            "terms.docx",
        )

        self.assertEqual(result.columns, ("word", "definition"))
        self.assertEqual(result.filename, "terms.csv")
        self.assertEqual(result.row_count, 2)
        rows = decode_tsv(result.content)
        self.assertEqual(rows[0], {"word": "abate", "definition": "become less intense"})
        self.assertEqual(rows[1]["definition"], "active and energetic")
        self.assertTrue(result.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"\r\n", result.content)

    def test_adapts_to_image_column_and_emits_webp(self):
        result = convert_dictionary_docx_to_csv(make_docx(True), "terms.docx")

        self.assertEqual(result.columns, ("word", "definition", "image"))
        rows = decode_tsv(result.content)
        self.assertTrue(rows[0]["image"].startswith("data:image/webp;base64,"))
        payload = rows[0]["image"].split(",", 1)[1]
        with Image.open(io.BytesIO(base64.b64decode(payload))) as image:
            self.assertEqual(image.format, "WEBP")
        self.assertEqual(rows[1]["image"], "")

    def test_ignores_formatting_list_markers_and_nested_table(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "word"
        table.cell(0, 1).text = "definition"
        table.cell(1, 0).text = "plain"

        definition = table.cell(1, 1)
        paragraph = definition.paragraphs[0]
        paragraph.style = "List Bullet"
        run = paragraph.add_run("visible definition")
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        nested = definition.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "nested content must not leak"

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        rows = decode_tsv(convert_dictionary_docx_to_csv(stream).content)
        self.assertEqual(rows[0]["definition"], "visible definition")
        self.assertNotIn("nested", rows[0]["definition"])

    def test_rejects_document_without_required_table(self):
        document = Document()
        document.add_paragraph("Not a dictionary table")
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        with self.assertRaisesRegex(ConversionError, "word.*definition"):
            convert_dictionary_docx_to_csv(stream)

    def test_rejects_invalid_docx(self):
        with self.assertRaises(ConversionError):
            convert_dictionary_docx_to_csv(io.BytesIO(b"not a docx"))


if __name__ == "__main__":
    unittest.main()
