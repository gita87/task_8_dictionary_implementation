"""Framework-independent DOCX dictionary conversion engine."""

from __future__ import annotations

import base64
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Callable, Iterable, Mapping, Protocol, TypeAlias

from docx import Document
from docx.document import Document as DocumentType
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table, _Cell
from PIL import Image, ImageOps, UnidentifiedImageError

from qaos_common import CancellationToken, ProcessingLimits
from qaos_common.progress import ProgressEvent as CommonProgressEvent
from qaos_common.limits import DEFAULT_LIMITS
from qaos_common.errors import QAOSCommonError, redact_sensitive
from qaos_common.csvio import DictionaryCSVProfile, DictionaryCSVWriter
from qaos_common.schemas import (
    DICTIONARY_COLUMNS, DICTIONARY_SCHEMA_VERSION, MISSING_DICTIONARY_IMAGE,
    validate_dictionary_headers, validate_dictionary_row,
)

REQUIRED_HEADERS = ("word", "definition")
OPTIONAL_HEADERS = ("image",)
GENERATED_HEADERS = ("unique_id",)
MISSING_IMAGE_VALUE = MISSING_DICTIONARY_IMAGE
WEBP_DATA_URI_PREFIX = "data:image/webp;base64,"
DEFAULT_IMAGE_QUALITY = 75
MAX_UPLOAD_BYTES = DEFAULT_LIMITS.max_upload_bytes
MAX_CELL_BYTES = DEFAULT_LIMITS.max_cell_bytes
MAX_OUTPUT_BYTES = DEFAULT_LIMITS.max_output_bytes
MAX_UNCOMPRESSED_DOCX_BYTES = MAX_OUTPUT_BYTES
MAX_DOCX_MEMBERS = 10_000

_ALLOWED_SOURCE_MIME_TYPES = frozenset({
    "image/bmp", "image/gif", "image/jpeg", "image/png", "image/tiff", "image/webp",
})
_ALLOWED_PIL_FORMATS = frozenset({"BMP", "GIF", "JPEG", "PNG", "TIFF", "WEBP"})
_MIME_TO_PIL_FORMAT = {
    "image/bmp": "BMP",
    "image/gif": "GIF",
    "image/jpeg": "JPEG",
    "image/png": "PNG",
    "image/tiff": "TIFF",
    "image/webp": "WEBP",
}
_WHITESPACE = re.compile(r"\s+")


@dataclass(frozen=True)
class Diagnostic:
    code: str
    message: str
    severity: str = "error"
    context: Mapping[str, object] = field(default_factory=dict)

    def as_dict(self) -> dict[str, object]:
        return redact_sensitive({"code": self.code, "message": self.message,
                "severity": self.severity, "context": dict(self.context)})


class ConversionError(QAOSCommonError, ValueError):
    """Conversion failure carrying a stable structured diagnostic."""

    def __init__(self, diagnostic: Diagnostic | str, *,
                 code: str = "conversion_failed",
                 context: Mapping[str, object] | None = None) -> None:
        if isinstance(diagnostic, str):
            diagnostic = Diagnostic(code, diagnostic, context=context or {})
        super().__init__(diagnostic.message, code=diagnostic.code,
                         details=diagnostic.context)
        self.diagnostic = Diagnostic(self.code, self.message, diagnostic.severity,
                                     self.details)

    def __str__(self) -> str:
        return self.message

    def as_dict(self) -> dict[str, object]:
        return self.diagnostic.as_dict()


@dataclass(frozen=True)
class ProgressEvent:
    stage: str
    completed: int
    total: int | None
    message: str

    def to_common(self) -> CommonProgressEvent:
        stages = {"read": "reading", "validate": "validating", "extract": "parsing",
                  "serialize": "writing", "complete": "completed"}
        return CommonProgressEvent(stages[self.stage], self.completed, self.total,
                                   self.message)


@dataclass(frozen=True)
class Checkpoint:
    schema_version: str
    stage: str
    rows_processed: int
    total_rows: int | None


class _EventLike(Protocol):
    def is_set(self) -> bool: ...


InputSource: TypeAlias = str | Path | bytes | bytearray | memoryview | BinaryIO
ProgressCallback: TypeAlias = Callable[[ProgressEvent], None]
CheckpointCallback: TypeAlias = Callable[[Checkpoint], None]
Cancellation: TypeAlias = CancellationToken | _EventLike | Callable[[], bool]


@dataclass(frozen=True)
class ConversionResult:
    content: bytes
    filename: str
    row_count: int
    columns: tuple[str, ...]
    schema_version: str = DICTIONARY_SCHEMA_VERSION
    diagnostics: tuple[Diagnostic, ...] = ()


def _error(code: str, message: str, **context: object) -> ConversionError:
    return ConversionError(Diagnostic(code, message, context=context))


def _cancelled(cancellation: Cancellation | None) -> bool:
    if cancellation is None:
        return False
    if isinstance(cancellation, CancellationToken):
        return cancellation.is_cancelled()
    if callable(cancellation):
        return bool(cancellation())
    return bool(cancellation.is_set())


def _check_cancelled(cancellation: Cancellation | None) -> None:
    if _cancelled(cancellation):
        raise _error("conversion_cancelled", "The conversion was cancelled.")


def _progress(callback: ProgressCallback | None, stage: str, completed: int,
              total: int | None, message: str) -> None:
    if callback is not None:
        callback(ProgressEvent(stage, completed, total, message))


def normalize_header(value: str) -> str:
    return _WHITESPACE.sub(" ", value).strip().casefold()


def plain_cell_text(cell: _Cell) -> str:
    paragraphs = (_WHITESPACE.sub(" ", p.text).strip() for p in cell.paragraphs)
    return " ".join(text for text in paragraphs if text).strip()


def _table_headers(table: Table) -> list[str]:
    if not table.rows:
        return []
    return [normalize_header(plain_cell_text(cell)) for cell in table.rows[0].cells]


def find_dictionary_table(tables: Iterable[Table]) -> tuple[Table, list[str]]:
    required = set(REQUIRED_HEADERS)
    candidates: list[tuple[int, int, Table, list[str]]] = []
    for order, table in enumerate(tables):
        headers = _table_headers(table)
        header_set = set(headers)
        if required.issubset(header_set):
            score = sum(header in header_set for header in OPTIONAL_HEADERS)
            candidates.append((score, -order, table, headers))
    if not candidates:
        raise _error("dictionary_table_not_found",
                     "No table was found with the required 'word' and 'definition' headers.",
                     required_headers=list(REQUIRED_HEADERS))
    _, _, table, headers = max(candidates, key=lambda item: (item[0], item[1]))
    return table, headers


def _first_image_blob(cell: _Cell) -> tuple[bytes, str] | None:
    for blip in cell._tc.xpath(".//a:blip"):
        relationship_id = blip.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        if not relationship_id:
            continue
        image_part = cell.part.related_parts.get(relationship_id)
        if image_part is not None and hasattr(image_part, "blob"):
            mime = getattr(image_part, "content_type", "application/octet-stream")
            return image_part.blob, mime.casefold()
    return None


def _validate_cell_size(value: str, column: str, row_number: int) -> None:
    size = len(value.encode("utf-8"))
    if size > MAX_CELL_BYTES:
        raise _error("cell_too_large",
                     f"Dictionary table row {row_number} column '{column}' exceeds the cell size limit.",
                     row=row_number, column=column, size_bytes=size,
                     limit_bytes=MAX_CELL_BYTES)


def image_to_data_uri(cell: _Cell, quality: int = DEFAULT_IMAGE_QUALITY) -> str:
    image_data = _first_image_blob(cell)
    if image_data is None:
        return MISSING_IMAGE_VALUE
    blob, content_type = image_data
    if content_type not in _ALLOWED_SOURCE_MIME_TYPES:
        raise _error("unsupported_image_mime",
                     "An embedded image uses an unsupported MIME type.",
                     mime_type=content_type,
                     allowed_mime_types=sorted(_ALLOWED_SOURCE_MIME_TYPES))
    if len(blob) > MAX_CELL_BYTES:
        raise _error("image_too_large", "An embedded image exceeds the image size limit.",
                     size_bytes=len(blob), limit_bytes=MAX_CELL_BYTES)
    try:
        with Image.open(io.BytesIO(blob)) as source:
            if source.format not in _ALLOWED_PIL_FORMATS:
                raise _error("unsupported_image_format",
                             "An embedded image uses an unsupported image format.",
                             detected_format=source.format)
            if source.format != _MIME_TO_PIL_FORMAT[content_type]:
                raise _error(
                    "image_mime_mismatch",
                    "An embedded image MIME type does not match its content.",
                    mime_type=content_type,
                    detected_format=source.format,
                )
            image = ImageOps.exif_transpose(source)
            image.load()
            if image.mode not in ("RGB", "RGBA"):
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            output = io.BytesIO()
            image.save(output, format="WEBP", quality=quality, method=6)
            value = WEBP_DATA_URI_PREFIX + base64.b64encode(output.getvalue()).decode("ascii")
            _validate_cell_size(value, "image", 0)
            return value
    except ConversionError:
        raise
    except (Image.DecompressionBombError, UnidentifiedImageError, OSError, ValueError) as error:
        raise _error("image_conversion_failed",
                     "An embedded image could not be converted to the required WebP format.",
                     mime_type=content_type) from error


def extract_rows(document: DocumentType, image_quality: int = DEFAULT_IMAGE_QUALITY, *,
                 progress_callback: ProgressCallback | None = None,
                 cancellation: Cancellation | None = None,
                 checkpoint_callback: CheckpointCallback | None = None,
                 checkpoint_interval: int = 100) -> tuple[tuple[str, ...], list[dict[str, str]]]:
    table, headers = find_dictionary_table(document.tables)
    indexes = {name: headers.index(name) for name in REQUIRED_HEADERS}
    has_image = "image" in headers
    if has_image:
        indexes["image"] = headers.index("image")
    columns = DICTIONARY_COLUMNS
    rows: list[dict[str, str]] = []
    source_rows = table.rows[1:]
    total = len(source_rows)
    for processed, table_row in enumerate(source_rows, start=1):
        _check_cancelled(cancellation)
        row_number = processed + 1
        cells = table_row.cells
        row: dict[str, str] = {}
        for column in REQUIRED_HEADERS:
            index = indexes[column]
            row[column] = plain_cell_text(cells[index]) if index < len(cells) else ""
            _validate_cell_size(row[column], column, row_number)
        if has_image:
            image_index = indexes["image"]
            row["image"] = (image_to_data_uri(cells[image_index], image_quality)
                            if image_index < len(cells) else MISSING_IMAGE_VALUE)
        else:
            row["image"] = MISSING_IMAGE_VALUE
        _validate_cell_size(row["image"], "image", row_number)
        _progress(progress_callback, "extract", processed, total, "Extracting rows")
        if checkpoint_callback is not None and (
                processed % checkpoint_interval == 0 or processed == total):
            checkpoint_callback(Checkpoint(DICTIONARY_SCHEMA_VERSION, "extract", processed, total))
        if not any(row[column].strip() for column in REQUIRED_HEADERS):
            continue
        missing = [column for column in REQUIRED_HEADERS if not row[column].strip()]
        if missing:
            raise _error("required_cell_empty",
                         f"Dictionary table row {row_number} has empty required cell(s): {', '.join(missing)}.",
                         row=row_number, columns=missing)
        row["unique_id"] = f"{len(rows) + 1:04d}"
        rows.append(row)
    return columns, rows


def rows_to_csv(columns: tuple[str, ...], rows: list[dict[str, str]], *,
                cancellation: Cancellation | None = None) -> bytes:
    """Serialize with the shared dictionary profile; retain legacy subset exports."""
    output = io.BytesIO()
    limits = ProcessingLimits(
        max_upload_bytes=MAX_UPLOAD_BYTES, max_cell_bytes=MAX_CELL_BYTES,
        max_output_bytes=MAX_OUTPUT_BYTES,
        max_image_bytes=min(DEFAULT_LIMITS.max_image_bytes, MAX_CELL_BYTES))
    canonical = set(DICTIONARY_COLUMNS).issubset(columns)
    seen_ids: set[str] = set()
    try:
        if canonical:
            validate_dictionary_headers(columns)
        with DictionaryCSVWriter(output, columns=columns,
                                 profile=DictionaryCSVProfile(line_ending="\r\n"),
                                 limits=limits) as writer:
            for row_number, row in enumerate(rows, start=2):
                _check_cancelled(cancellation)
                exported = dict(row)
                if canonical:
                    validate_dictionary_row(exported, row_number=row_number,
                                            seen_ids=seen_ids)
                    if not exported["definition"].strip():
                        raise _error("required_cell_empty", "Dictionary definition is empty.",
                                     row=row_number, columns=["definition"])
                if "unique_id" in columns:
                    exported["unique_id"] = f"'{row['unique_id']}"
                writer.write_row(exported)
    except ConversionError:
        raise
    except QAOSCommonError as error:
        code = {"CSV_OUTPUT_TOO_LARGE": "output_too_large",
                "CSV_CELL_TOO_LARGE": "cell_too_large"}.get(error.code, error.code)
        converted = ConversionError(error.message, code=code, context=error.details)
        converted.stage = error.stage
        raise converted from error
    return output.getvalue()


def _output_filename(input_filename: str) -> str:
    return f"{Path(input_filename).stem.strip() or 'dictionary'}.csv"


def _read_source(source: InputSource, input_filename: str | None) -> tuple[bytes, str]:
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = input_filename or path.name
        try:
            with path.open("rb") as stream:
                payload = stream.read(MAX_UPLOAD_BYTES + 1)
        except OSError as error:
            raise _error("input_read_failed", "The DOCX file could not be read.", path=str(path)) from error
    elif isinstance(source, (bytes, bytearray, memoryview)):
        payload, filename = bytes(source), input_filename or "dictionary.docx"
    else:
        filename = input_filename or Path(getattr(source, "name", "dictionary.docx")).name
        try:
            payload = source.read(MAX_UPLOAD_BYTES + 1)
        except (AttributeError, OSError, ValueError) as error:
            raise _error("input_read_failed", "The DOCX stream could not be read.") from error
    if not payload:
        raise _error("input_empty", "The DOCX file is empty.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise _error("input_too_large", "The DOCX exceeds the upload size limit.",
                     size_bytes=len(payload), limit_bytes=MAX_UPLOAD_BYTES)
    return payload, filename


def _validated_docx_buffer(payload: bytes) -> io.BytesIO:
    buffer = io.BytesIO(payload)
    try:
        with zipfile.ZipFile(buffer) as package:
            members = package.infolist()
            names = {member.filename for member in members}
            expanded_size = sum(member.file_size for member in members)
            if len(members) > MAX_DOCX_MEMBERS:
                raise _error("docx_too_many_parts", "The DOCX contains too many internal parts.")
            if expanded_size > MAX_UNCOMPRESSED_DOCX_BYTES:
                raise _error("docx_uncompressed_too_large",
                             "The DOCX is too large after decompression.",
                             size_bytes=expanded_size,
                             limit_bytes=MAX_UNCOMPRESSED_DOCX_BYTES)
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise _error("invalid_docx_package", "The file is not a valid DOCX package.")
    except zipfile.BadZipFile as error:
        raise _error("invalid_docx_zip", "The file could not be read as a valid DOCX.") from error
    buffer.seek(0)
    return buffer


def convert_dictionary_docx(source: InputSource, input_filename: str | None = None,
                            image_quality: int = DEFAULT_IMAGE_QUALITY, *,
                            progress_callback: ProgressCallback | None = None,
                            cancellation: Cancellation | None = None,
                            checkpoint_callback: CheckpointCallback | None = None,
                            checkpoint_interval: int = 100,
                            common_progress_callback: Callable[[CommonProgressEvent], None] | None = None,
                            ) -> ConversionResult:
    """Convert a DOCX path, byte string, or binary stream to dictionary/1.0 CSV."""
    if not 1 <= image_quality <= 100:
        raise ValueError("image_quality must be between 1 and 100.")
    if checkpoint_interval < 1:
        raise ValueError("checkpoint_interval must be at least 1.")
    legacy_callback = progress_callback
    if common_progress_callback is not None:
        def emit(event: ProgressEvent) -> None:
            if legacy_callback is not None:
                legacy_callback(event)
            common_progress_callback(event.to_common())
        progress_callback = emit
    _check_cancelled(cancellation)
    _progress(progress_callback, "read", 0, None, "Reading input")
    payload, filename = _read_source(source, input_filename)
    _check_cancelled(cancellation)
    _progress(progress_callback, "validate", 1, 1, "Validating DOCX")
    try:
        document = Document(_validated_docx_buffer(payload))
    except ConversionError:
        raise
    except (PackageNotFoundError, KeyError, OSError, ValueError) as error:
        raise _error("invalid_docx", "The file could not be read as a valid DOCX.") from error
    except Exception as error:
        raise _error("unsupported_docx_structure",
                     "The DOCX structure is corrupted or unsupported.") from error
    _check_cancelled(cancellation)
    columns, rows = extract_rows(
        document, image_quality=image_quality, progress_callback=progress_callback,
        cancellation=cancellation, checkpoint_callback=checkpoint_callback,
        checkpoint_interval=checkpoint_interval)
    _check_cancelled(cancellation)
    _progress(progress_callback, "serialize", 0, 1, "Serializing CSV")
    content = rows_to_csv(columns, rows, cancellation=cancellation)
    _check_cancelled(cancellation)
    _progress(progress_callback, "complete", 1, 1, "Conversion complete")
    return ConversionResult(content, _output_filename(filename), len(rows), columns)


def convert_dictionary_docx_to_csv(source: InputSource,
                                   input_filename: str = "dictionary.docx",
                                   image_quality: int = DEFAULT_IMAGE_QUALITY,
                                   **kwargs: object) -> ConversionResult:
    """Backward-compatible alias for :func:`convert_dictionary_docx`."""
    return convert_dictionary_docx(source, input_filename, image_quality, **kwargs)
