"""Core DOCX dictionary-table to tab-delimited CSV conversion engine."""

from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Iterable

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table, _Cell
from PIL import Image, ImageOps, UnidentifiedImageError


REQUIRED_HEADERS = ("word", "definition")
OPTIONAL_HEADERS = ("image",)
GENERATED_HEADERS = ("unique_id",)
WEBP_DATA_URI_PREFIX = "data:image/webp;base64,"
DEFAULT_IMAGE_QUALITY = 75
MAX_UNCOMPRESSED_DOCX_BYTES = 250 * 1024 * 1024
MAX_DOCX_MEMBERS = 10_000
_WHITESPACE = re.compile(r"\s+")


class ConversionError(ValueError):
    """A safe, user-facing conversion failure."""


@dataclass(frozen=True)
class ConversionResult:
    content: bytes
    filename: str
    row_count: int
    columns: tuple[str, ...]


def normalize_header(value: str) -> str:
    """Normalize headers for matching without changing CSV column names."""
    return _WHITESPACE.sub(" ", value).strip().casefold()


def plain_cell_text(cell: _Cell) -> str:
    """Extract visible paragraph text while ignoring formatting and nested tables."""
    paragraphs = (
        _WHITESPACE.sub(" ", paragraph.text).strip()
        for paragraph in cell.paragraphs
    )
    return " ".join(text for text in paragraphs if text).strip()


def _table_headers(table: Table) -> list[str]:
    if not table.rows:
        return []
    return [normalize_header(plain_cell_text(cell)) for cell in table.rows[0].cells]


def find_dictionary_table(tables: Iterable[Table]) -> tuple[Table, list[str]]:
    """Select the best top-level table containing the required dictionary headers."""
    required = set(REQUIRED_HEADERS)
    candidates: list[tuple[int, int, Table, list[str]]] = []

    for order, table in enumerate(tables):
        headers = _table_headers(table)
        header_set = set(headers)
        if required.issubset(header_set):
            preferred_score = sum(header in header_set for header in OPTIONAL_HEADERS)
            candidates.append((preferred_score, -order, table, headers))

    if not candidates:
        raise ConversionError(
            "No table was found with the required 'word' and 'definition' headers."
        )

    _, _, table, headers = max(candidates, key=lambda item: (item[0], item[1]))
    return table, headers


def _first_image_blob(cell: _Cell) -> tuple[bytes, str] | None:
    """Return the first embedded image in a cell, if present."""
    for blip in cell._tc.xpath(".//a:blip"):
        relationship_id = blip.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        if not relationship_id:
            continue

        image_part = cell.part.related_parts.get(relationship_id)
        if image_part is None or not hasattr(image_part, "blob"):
            continue

        content_type = getattr(image_part, "content_type", "application/octet-stream")
        return image_part.blob, content_type

    return None


def image_to_data_uri(cell: _Cell, quality: int = DEFAULT_IMAGE_QUALITY) -> str:
    """Convert a cell's first embedded image to the required WebP data URI."""
    image_data = _first_image_blob(cell)
    if image_data is None:
        return ""

    blob, _content_type = image_data
    try:
        with Image.open(io.BytesIO(blob)) as source:
            image = ImageOps.exif_transpose(source)
            image.load()
            if image.mode not in ("RGB", "RGBA"):
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")

            output = io.BytesIO()
            image.save(output, format="WEBP", quality=quality, method=6)
            encoded = base64.b64encode(output.getvalue()).decode("ascii")
            return f"{WEBP_DATA_URI_PREFIX}{encoded}"
    except (
        Image.DecompressionBombError,
        UnidentifiedImageError,
        OSError,
        ValueError,
    ) as error:
        raise ConversionError(
            "An embedded image could not be converted to the required WebP format."
        ) from error


def extract_rows(document: Document, image_quality: int = DEFAULT_IMAGE_QUALITY):
    """Extract the target columns and data rows from a DOCX document."""
    table, headers = find_dictionary_table(document.tables)

    column_indexes = {name: headers.index(name) for name in REQUIRED_HEADERS}
    has_image = "image" in headers
    if has_image:
        column_indexes["image"] = headers.index("image")

    columns = tuple(
        GENERATED_HEADERS
        + REQUIRED_HEADERS
        + (OPTIONAL_HEADERS if has_image else ())
    )
    rows: list[dict[str, str]] = []

    for row_number, table_row in enumerate(table.rows[1:], start=2):
        cells = table_row.cells
        row: dict[str, str] = {}

        for column in REQUIRED_HEADERS:
            index = column_indexes[column]
            row[column] = plain_cell_text(cells[index]) if index < len(cells) else ""

        if has_image:
            image_index = column_indexes["image"]
            row["image"] = (
                image_to_data_uri(cells[image_index], image_quality)
                if image_index < len(cells)
                else ""
            )

        if not any(value.strip() for value in row.values()):
            continue

        missing_required = [
            column for column in REQUIRED_HEADERS if not row[column].strip()
        ]
        if missing_required:
            missing = ", ".join(missing_required)
            raise ConversionError(
                f"Dictionary table row {row_number} has empty required "
                f"cell(s): {missing}."
            )

        row["unique_id"] = f"{len(rows) + 1:04d}"
        rows.append(row)

    return columns, rows


def rows_to_csv(columns: tuple[str, ...], rows: list[dict[str, str]]) -> bytes:
    """Serialize rows using the tab-delimited UTF-8 BOM format from the reference."""
    output = io.StringIO(newline="")
    writer = csv.DictWriter(
        output,
        fieldnames=columns,
        delimiter="\t",
        lineterminator="\r\n",
        quoting=csv.QUOTE_MINIMAL,
        quotechar='"',
        doublequote=True,
    )
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def _output_filename(input_filename: str) -> str:
    stem = Path(input_filename).stem.strip() or "dictionary"
    return f"{stem}.csv"


def _validated_docx_buffer(source: BinaryIO) -> io.BytesIO:
    """Copy and validate an uploaded DOCX ZIP before python-docx parses it."""
    try:
        payload = source.read()
    except (AttributeError, OSError, ValueError) as error:
        raise ConversionError("The uploaded DOCX could not be read.") from error

    if not payload:
        raise ConversionError("The DOCX file is empty.")

    buffer = io.BytesIO(payload)
    try:
        with zipfile.ZipFile(buffer) as package:
            members = package.infolist()
            names = {member.filename for member in members}
            uncompressed_size = sum(member.file_size for member in members)

            if len(members) > MAX_DOCX_MEMBERS:
                raise ConversionError("The DOCX contains too many internal parts.")
            if uncompressed_size > MAX_UNCOMPRESSED_DOCX_BYTES:
                raise ConversionError("The DOCX is too large after decompression.")
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ConversionError("The file is not a valid DOCX package.")
    except zipfile.BadZipFile as error:
        raise ConversionError("The file could not be read as a valid DOCX.") from error

    buffer.seek(0)
    return buffer


def convert_dictionary_docx_to_csv(
    source: BinaryIO,
    input_filename: str = "dictionary.docx",
    image_quality: int = DEFAULT_IMAGE_QUALITY,
) -> ConversionResult:
    """Convert an uploaded DOCX stream into a downloadable CSV result."""
    if not 1 <= image_quality <= 100:
        raise ValueError("image_quality must be between 1 and 100.")

    try:
        document = Document(_validated_docx_buffer(source))
    except ConversionError:
        raise
    except (PackageNotFoundError, KeyError, OSError, ValueError) as error:
        raise ConversionError(
            "The file could not be read as a valid DOCX."
        ) from error
    except Exception as error:
        # python-docx can surface malformed ZIP/XML failures through several types.
        raise ConversionError(
            "The DOCX structure is corrupted or unsupported."
        ) from error

    columns, rows = extract_rows(document, image_quality=image_quality)
    return ConversionResult(
        content=rows_to_csv(columns, rows),
        filename=_output_filename(input_filename),
        row_count=len(rows),
        columns=columns,
    )
