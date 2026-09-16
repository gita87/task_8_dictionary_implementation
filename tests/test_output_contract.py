from __future__ import annotations

import csv
import io
import unittest

from docx import Document

from dict_docx_to_csv import (
    ConversionError,
    WEBP_DATA_URI_PREFIX,
    convert_dictionary_docx_to_csv,
)
from tests.test_dict_docx_to_csv import make_docx


def save_document(document: Document) -> io.BytesIO:
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


class OutputContractTests(unittest.TestCase):
    def test_emits_exact_bom_tab_crlf_and_quote_contract(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "word"
        table.cell(0, 1).text = "definition"
        table.cell(1, 0).text = 'alpha\t"beta"'
        table.cell(1, 1).text = "first\nsecond"

        result = convert_dictionary_docx_to_csv(save_document(document))

        self.assertEqual(
            result.content,
            (
                b'\xef\xbb\xbfunique_id\tword\tdefinition\r\n'
                b"'0001\t\"alpha \"\"beta\"\"\"\tfirst second\r\n"
            ),
        )
        self.assertNotIn(b"\n", result.content.replace(b"\r\n", b""))

    def test_emits_exact_lowercase_column_order(self):
        document = Document()
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "IMAGE"
        table.cell(0, 1).text = "DEFINITION"
        table.cell(0, 2).text = "WORD"
        table.cell(1, 1).text = "a definition"
        table.cell(1, 2).text = "a word"

        result = convert_dictionary_docx_to_csv(save_document(document))

        self.assertEqual(
            result.columns, ("unique_id", "word", "definition", "image")
        )
        header = result.content.decode("utf-8-sig").split("\r\n", 1)[0]
        self.assertEqual(header, "unique_id\tword\tdefinition\timage")

    def test_image_value_is_a_raw_webp_data_uri_without_html(self):
        result = convert_dictionary_docx_to_csv(make_docx(True), "terms.docx")
        text = result.content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text), delimiter="\t"))
        image_value = rows[0]["image"]

        self.assertTrue(image_value.startswith(WEBP_DATA_URI_PREFIX))
        self.assertNotIn("<img", image_value.casefold())
        self.assertNotIn("[src]", image_value.casefold())
        self.assertNotIn("data:image/png", image_value.casefold())
        self.assertNotIn("data:image/jpeg", image_value.casefold())

    def test_rejects_a_partially_populated_required_row(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "word"
        table.cell(0, 1).text = "definition"
        table.cell(1, 0).text = "missing definition"

        with self.assertRaisesRegex(
            ConversionError,
            r"row 2.*definition",
        ):
            convert_dictionary_docx_to_csv(save_document(document))

    def test_ignores_a_fully_empty_table_row(self):
        document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "word"
        table.cell(0, 1).text = "definition"
        table.cell(2, 0).text = "valid"
        table.cell(2, 1).text = "valid definition"

        result = convert_dictionary_docx_to_csv(save_document(document))

        self.assertEqual(result.row_count, 1)
        rows = list(
            csv.DictReader(
                io.StringIO(result.content.decode("utf-8-sig")), delimiter="\t"
            )
        )
        self.assertEqual(rows[0]["unique_id"], "'0001")


if __name__ == "__main__":
    unittest.main()
